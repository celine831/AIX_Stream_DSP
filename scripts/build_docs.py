"""Generate the two Chinese engineering Word deliverables from measured reports."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from PIL import Image,ImageDraw,ImageFont
from collect_results import collect,ROOT

QA=ROOT/'tmp/doc_qa'
QA.mkdir(parents=True,exist_ok=True)

def font(style,name='宋体',size=11):
    style.font.name=name;style.font.size=Pt(size)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),name)

def new(title,intro):
    d=Document();sec=d.sections[0]
    sec.page_width=Inches(8.2677);sec.page_height=Inches(11.6929)
    sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.65)
    sec.left_margin=Inches(.8);sec.right_margin=Inches(.8)
    normal=d.styles['Normal'];font(normal)
    normal.paragraph_format.space_after=Pt(7);normal.paragraph_format.line_spacing=1.12
    for name,size in [('Title',23),('Heading 1',16),('Heading 2',12)]:
        font(d.styles[name],'微软雅黑',size);d.styles[name].font.color.rgb=RGBColor(0,0,0)
        d.styles[name].paragraph_format.space_before=Pt(12)
        d.styles[name].paragraph_format.space_after=Pt(8)
        for border in d.styles[name].element.xpath('./w:pPr/w:pBdr'):
            border.getparent().remove(border)
    d.add_paragraph(title,'Title');d.add_paragraph('版本 V1.1    目标器件 EP4CE10F17C8    日期 2026年10月13日')
    d.add_paragraph(intro)
    p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run('第 ');field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');p._p.append(field);p.add_run(' 页')
    d.core_properties.title=title;d.core_properties.author='张佳凝';d.core_properties.last_modified_by='张佳凝'
    d.core_properties.subject='参数化 AXI4 Stream 定点 FIR 加速器';d.core_properties.comments='';d.core_properties.revision=2
    return d

def page(d,title):
    d.add_page_break()
    if title.startswith('附录'):
        d.add_heading(title,1)
    else:
        d._chapter=getattr(d,'_chapter',1)+1
        d.add_heading(str(d._chapter)+' '+title,1)

def h(d,title,text):
    d.add_heading(title,2);d.add_paragraph(text)

def p(d,text):d.add_paragraph(text)

def table(d,headers,rows,widths=None):
    t=d.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.style=None;t.autofit=False
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        element=OxmlElement('w:'+edge)
        for key,value in [('val','single'),('sz','4'),('color','D9D9D9')]:element.set(qn('w:'+key),value)
        borders.append(element)
    t._tbl.tblPr.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for edge in ['top','bottom','left','right']:
        element=OxmlElement('w:'+edge);element.set(qn('w:w'),'70');element.set(qn('w:type'),'dxa');margins.append(element)
    t._tbl.tblPr.append(margins)
    # Neutral engineering table rather than a colored template theme.
    for cell,value in zip(t.rows[0].cells,headers):cell.text=value
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        for cell,value in zip(t.add_row().cells,row):cell.text=str(value)
    if widths:
        widths=[width*6.6/sum(widths) for width in widths]
        for col,w in zip(t.columns,widths):col.width=Inches(w)
        for row in t.rows:
            for cell,w in zip(row.cells,widths):cell.width=Inches(w)
    for ri,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        for cell in row.cells:
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri==0 or ri%2==1:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EDEDED' if ri==0 else 'F8F8F8');cell._tc.get_or_add_tcPr().append(shade)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(4);para.paragraph_format.space_before=Pt(4)
                for run in para.runs:
                    run.font.size=Pt(10.5);run.bold=(ri==0);run.font.color.rgb=RGBColor(0,0,0)
    d.add_paragraph('')

def equation(d):
    para=d.add_paragraph();math=OxmlElement('m:oMath')
    def run(parent,text):
        r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=text;r.append(t);parent.append(r)
    run(math,'y[n] = ')
    nary=OxmlElement('m:nary');pr=OxmlElement('m:naryPr');c=OxmlElement('m:chr');c.set(qn('m:val'),'∑');pr.append(c)
    loc=OxmlElement('m:limLoc');loc.set(qn('m:val'),'undOvr');pr.append(loc);nary.append(pr)
    sub=OxmlElement('m:sub');run(sub,'k = 0');nary.append(sub)
    sup=OxmlElement('m:sup');run(sup,'TAPS − 1');nary.append(sup)
    e=OxmlElement('m:e');run(e,'c[k] x[n − k]');nary.append(e);math.append(nary)
    para._p.append(math)

def references(d):
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    para=d.add_paragraph('参考资料  ')
    for label,url in [('ARM AXI4 Stream 规范','https://documentation-service.arm.com/static/642583d7314e245d086bc8c9'),
                      ('Altera 双时钟 FIFO 约束','https://docs.altera.com/r/docs/683082/25.1/quartus-prime-pro-edition-user-guide/dual-clock-fifo-timing-constraints')]:
        link=OxmlElement('w:hyperlink');link.set(qn('r:id'),d.part.relate_to(url,RT.HYPERLINK,is_external=True))
        run=OxmlElement('w:r');text=OxmlElement('w:t');text.text=label;run.append(text);link.append(run);para._p.append(link);para.add_run('  ')
    para.add_run('Quartus II 13.0.1 时序分析命令参考')

def diagram():
    img=Image.new('RGB',(1600,580),'white');draw=ImageDraw.Draw(img)
    f=ImageFont.truetype('C:/Windows/Fonts/msyh.ttc',26)
    small=ImageFont.truetype('C:/Windows/Fonts/msyh.ttc',23)
    labels=['输入 FIFO\n同步或异步','输入弹性\n寄存器','并行乘法器\n乘积寄存','流水化\n加法树','定点量化\n输出寄存','输出 FIFO\n弹性寄存器']
    for i,label in enumerate(labels):
        x=30+260*i;draw.rounded_rectangle((x,90,x+230,235),radius=8,fill='#f5f5f5',outline='#444444',width=3)
        draw.multiline_text((x+15,125),label,fill='#192c38',font=f,spacing=12)
        if i<5:draw.line((x+231,165,x+260,165),fill='#34566b',width=4);draw.polygon([(x+260,165),(x+246,157),(x+246,173)],fill='#34566b')
    draw.line((55,300,1530,300),fill='#678195',width=4)
    draw.polygon([(55,300),(70,292),(70,308)],fill='#678195')
    draw.text((620,315),'TREADY 逐级传播反压',fill='#222222',font=small)
    draw.rounded_rectangle((420,410,1150,510),radius=15,fill='#f4f4f4',outline='#7a838a',width=2)
    draw.text((440,440),'影子系数写入 → 暂停输入 → 排空 → 原子提交',font=f,fill='#222222')
    path=QA/'architecture.png';img.save(path);return path

def measured_rows(r):
    rows=[]
    for rev,v in r['revisions'].items():
        if rev=='board_selftest':continue
        z=v['resources'];sl=v['slow_85C_slack_ns']
        rows.append([rev,str(v['fmax_mhz']['core_clk'])+' MHz',z['Total logic elements'],z['Total registers'],z['Embedded Multiplier 9-bit elements'],sl['setup']])
    return rows

def planning(r):
    d=new('AXI4 Stream FIR DSP 加速器设计说明书','本项目实现面向连续采样数据的参数化定点 FIR 加速器。输入与输出采用 AXI4 Stream 握手传输，运算通道由并行乘法器、流水加法树和定点量化单元组成，支持反压、运行时系数更新及可选异步输入。本说明书规定系统架构、接口行为、数值精度、时钟复位和验证要求。')
    d.add_heading('1 项目概述与技术指标',1)
    table(d,['技术要求','设计方案','验收指标'],[
        ['位宽 阶数 系数可配置','位宽和抽头参数化；运行时影子系数表','11 组 FIR 参数测试'],
        ['有符号定点乘加','并行有符号乘法；完整精度加法树','Python 整数黄金结果'],
        ['AXI 握手和反压','输入 输出 FIFO 与弹性寄存器','停顿稳定性和连续吞吐断言'],
        ['可选 CDC','双时钟 Gray 指针 FIFO','快慢源时钟 仿真与 Gray skew 报告'],
        ['舍入 截位 饱和','3 种量化模式与饱和 回绕选择','122880 个穷举比对'],
        ['EDA 验证和分析','ModelSim 功能仿真；Quartus 综合及 STA','编译成功；规定频率下时序闭合']], [1.65,2.6,2.6])
    p(d,'系统由 7 个核心 RTL 模块、板级自检顶层和参考 ROM 组成。验证系统包含 4 个测试平台及独立整数参考模型。单时钟、双时钟输入和板级自检分别对应三个实现配置，使用同一套运算 RTL。')
    p(d,'目标器件为 Cyclone IV E EP4CE10F17C8。核心性能配置的工作频率为 75 MHz，CDC 源域为 50 MHz，板级自检采用 50 MHz 晶振。设计以数值一致性、协议行为、资源容量及全部时序模型满足约束为验收依据。')

    page(d,'总体架构与数值设计')
    d.add_picture(str(diagram()),width=Inches(6.8))
    p(d,'输入 FIFO 吸收突发数据，输入寄存器隔离存储读出与乘法路径。每个被 FIR 接收的样本与当前历史样本同时进入各抽头乘法器，乘积在平衡加法树逐级相加。量化后的输出进入 FIFO 和寄存器，从而保持反压期间的 AXI 数据稳定。')
    equation(d)
    table(d,['参数','默认值','含义'],[
        ['DATA_WIDTH COEFF_WIDTH OUT_WIDTH','16 16 16','样本 系数 输出有符号位宽'],
        ['TAPS SHIFT','8 15','8 抽头 7 阶；移除 15 个低位'],
        ['ROUND_MODE SATURATE','1 1','中点远离零；溢出饱和'],
        ['FRAME_RESET USE_ASYNC','0 0','跨包保持历史；默认单时钟'],
        ['FIFO_ADDR_WIDTH','4','输入 输出各 16 个存储位置']], [2.55,1.2,3.1])
    p(d,'默认乘积为 32 位，加法树有 3 级，累加器为 35 位。抽头数不是二次幂时补零到最近的二次幂。完整精度累加后才进行舍入和饱和，避免每一级丢失低位。输入若为整数单位，15 个小数位的系数配合 SHIFT 为 15 可保持输出单位。')

    page(d,'外部接口与系数更新规格')
    table(d,['信号','时钟域','协议'],[
        ['s_axis_tdata tvalid tready tlast','源域 或 aclk','valid 与 ready 同时为 1 才接收'],
        ['m_axis_tdata tvalid tready tlast','aclk','下游停顿时 valid 与 payload 保持'],
        ['cfg_we cfg_addr cfg_data','aclk','上升沿写影子表；地址范围 0 至 TAPS 减 1'],
        ['cfg_quiesce cfg_ready','aclk','停止新输入并等待所有已接收数据排空'],
        ['cfg_commit cfg_error','aclk','合法提交一次完成；错误组合指示'],
        ['aresetn','共同异步复位','异步拉低 各域同步释放']], [2.7,1.1,3.05])
    h(d,'数据总线与帧语义','AXI TDATA 总线向上补齐到字节宽度，非整字节数值只解释低 DATA_WIDTH 位，输出高位符号扩展。省略 TKEEP 和 TSTRB 表示所有字节有效。TLAST 原样随数据传播，不插入 FIR 尾部补零；FRAME_RESET 为 1 时末样本仍使用本帧历史，但其接收后历史清零。')
    h(d,'配置操作顺序','先逐项写 cfg_we、cfg_addr 和 cfg_data，再保持 cfg_quiesce 为 1 并等待 cfg_ready。下游继续拉 ready 排空。提交时 cfg_commit 拉高一拍且 cfg_we 为 0，然后撤销暂停。提交复制完整系数表并清空历史；未改写的影子系数保留原值。')
    h(d,'错误与初始状态','复位将两套系数表清零。非法系数地址不写入；未排空的 commit 不生效；同拍 write 与 commit 仅写影子表，提交被拒绝并给出 cfg_error。配置接口为简单同步控制接口，未实现 AXI4 Lite 总线。')

    page(d,'跨时钟复位与时序约束设计')
    h(d,'跨时钟数据通道','只有 USE_ASYNC 为 1 时输入 FIFO 使用源时钟写入和核心时钟读出。二进制指针用于寻址，Gray 指针通过两级同步器传递，带额外回绕位实现满空判断。数据存储和指针一起保证读取内容先写入再可见，不将每位 payload 单独同步。')
    h(d,'配置暂停确认','暂停请求从核心域同步到源域，源端停止新的握手，再将确认同步回核心域。确认返回后额外等待 3 个核心时钟形成 pause_settled，覆盖最后一次写指针同步与 empty 标志寄存器更新。cfg_ready 同时要求输入 FIFO、输入寄存器、FIR、输出 FIFO 和输出寄存器均空。')
    h(d,'共同复位契约','两个域必须共享 aresetn，复位会丢弃未完成数据。复位释放采用两级同步，并要求相关时钟运行。单独复位一个域会破坏 FIFO 指针对应关系，当前设计不支持该操作。RTL 仿真不模拟模拟电路的亚稳态；同步器结构和物理约束降低风险但不能提供形式上的绝对证明。')
    h(d,'跨域物理约束','Gray 总线的最大 skew 为 10 ns，小于本配置最短发起周期 13.333 ns。指针跨域路径同时设置最大延迟与最小延迟，暂停和确认的第一拍同步器单独设置时序例外。整体异步时钟组切断会屏蔽 Gray 总线检查，因此约束按具体路径分类设置。')
    p(d,'同步版本仅创建核心时钟。CDC 版本分别创建 13.333 ns 和 20 ns 时钟，所有常规输入 输出设置最大 2 ns、最小 0 ns 外部延迟预算；共同异步复位的外部路径明确切除。该预算用于核心工程验证，接入实际板卡后必须按系统时钟与外设重新计算。')

    page(d,'验证方案与验收指标')
    table(d,['验证层次','覆盖内容','通过判据'],[
        ['数值 黄金模型','冲激 极值 符号交替 常量 随机序列','输出及 TLAST 全部一致'],
        ['流协议','连续输入 随机空拍 80 拍长阻塞','无丢失 重复 乱序；停顿稳定'],
        ['参数组合','1 3 5 8 16 抽头 与多种位宽','11 个配置全部 PASS'],
        ['CDC','源时钟比核心快或慢 Gray 指针回绕','指针单比特变化 数据正确'],
        ['复位与配置','满 FIFO 复位 排空暂停 非法提交','无旧数据残留 错误被拒绝'],
        ['定点量化','4096 个输入乘 6 种模式乘 5 个 SHIFT','122880 项比对无错误'],
        ['综合和 STA','三个 revision 完整编译 全部 PVT 角落','工作时钟下无时序违例']], [1.2,3.1,2.55])
    h(d,'功能验证方法','独立参考模型依据输入接收顺序计算结果，输出计分板逐项比较数值和 TLAST，并检查接收数与输出数一致。源端停顿时保持数据，消费者制造随机反压和长阻塞。随机种子固定，确保相同输入可重复生成。')
    h(d,'实现验证方法','三个实现配置分别完成综合、布局布线及静态时序分析。验收要求编译成功，工作频率下的 setup、hold、recovery、removal 和脉宽检查通过，常规接口约束完整，跨域约束匹配实际路径。')
    h(d,'结果记录','仿真日志和 WLF 波形位于 sim/results，综合和时序报告位于 Qproject/reports。summary.json 记录各用例结果，verification_summary.json 汇总比较数量、资源占用和时序指标，用于对应源文件与验证记录。')

    page(d,'性能结果与适用范围')
    table(d,['配置','核心 Fmax','LE','寄存器','9位乘法资源','Setup ns'],measured_rows(r),[1.45,1.12,.6,.75,1.25,1.08])
    v=r['revisions']['accelerator_cdc']
    p(d,f"CDC 源时钟 Fmax 为 {v['fmax_mhz']['source_clk']} MHz。两种配置的工作核心时钟均为约 75 MHz；CDC 源为 50 MHz。完整编译成功，三个温度 电压时序模型均满足约束。no_clock、multiple_clock、no_input_delay 和 no_output_delay 均为 0；未约束路径为 0。")
    p(d,f"ModelSim 共 {r['simulation_tests']} 项测试通过，FIR 比对 {r['fir_comparisons']} 项，量化比对 {r['quantizer_comparisons']} 项。默认同步模式连续输出 260 拍，首个输入到输出握手间隔为 9 个理想核心时钟。Gray skew 20 条报告路径无违例，最差裕量为 {r['gray_skew_worst_slack_ns']} ns。")
    h(d,'频率选择依据','初始 100 MHz 约束下 Fmax 为 66.31 MHz，性能限制主要来自 FIFO 读出路径。增加接口寄存器、浅 FIFO 映射为逻辑存储并简化舍入后，同步与 CDC 核心的 100 MHz 实现结果分别为 83.40 MHz 和 87.57 MHz。最终工作频率规定为 75 MHz，并在该约束下重新布局布线，以上表格为最终实现结果。')
    h(d,'资源与板级配置','两个核心配置均使用 16 个 9 位嵌入乘法元素，占 46 个可用元素的约 35%，对应 8 路 16 位乘法。浅 FIFO 使用逻辑寄存器，M9K 存储位使用量为零。核心配置用于性能评估，外部 AXI 管脚自动分配；板级运行采用已绑定晶振、复位和 LED 的 board_selftest 配置。')
    h(d,'适用范围','当前数据通道支持规定范围内的参数化结构及固定总线精度。修改位宽、抽头、FIFO 深度或外部接口后，应重新进行功能与时序验证。50 MHz 单时钟板级自检和复位重测已通过；异步 CDC 的验收范围为仿真、同步器结构和物理约束，不包含双时钟实物试验。')
    board_page(d,r)
    d.save(ROOT/'doc/01_项目设计说明书.docx')

def board_page(d,r):
    page(d,'板级自检设计与实测结果')
    table(d,['接口','管脚','板上标识','功能'],[['sys_clk','E1','晶振','50 MHz 时钟'],['sys_rst_n','M15','RESET','低电平复位'],['led[0]','L7','D6','时钟心跳'],['led[1]','M6','D7','自检通过 PASS'],['led[2]','P3','D8','自检失败 FAIL'],['led[3]','N3','D9','运行状态 RUN']], [1.25,.65,.9,4.05])
    p(d,'以上六个管脚使用 3.3 V LVTTL；LED 低电平点亮，输出驱动配置为 4 mA。未使用的普通用户引脚在拟合报告中为输入三态带弱上拉。绑定放在 Qproject/board_pins.qsf，板级 revision 为 board_selftest。')
    h(d,'自检控制与比较','配置后通过上电计数产生初始化复位，控制状态依次为 BOOT、PAUSE、WRITE、COMMIT、RELEASE、RUN、DRAIN 和 DONE。四套 8 抽头系数各处理 64 个样本，共比较 256 项数值及 TLAST，包含极值、舍入、饱和和随机数据。接收端主动制造长短反压，源端和比较端分别按发送、接收计数寻址 ROM。错误或每组 1 秒超时均锁存 FAIL，全部比较完成且无错误才输出 PASS。')
    b=r['revisions']['board_selftest']
    p(d,f"板级配置在全部时序模型下满足 50 MHz 工作约束，报告 Fmax 为 {b['fmax_mhz']['board_clk']} MHz，使用 {b['resources']['Total logic elements']} 个 LE。自检测试平台两次运行共完成 512 项结果比较，另覆盖心跳和故障注入后的 FAIL 锁存。板级版本采用单时钟，CDC 由独立配置验证。")
    h(d,'实物验证记录','测试日期为 2026年10月13日，记录方式为现场人工观察。运行结果为 D6 持续闪烁、D7 常亮、D8 和 D9 熄灭。按住 RESET 时 D6、D7 熄灭，释放后恢复心跳闪烁和 PASS 常亮，其余用户灯保持熄灭。该结果表明 50 MHz 单时钟自检与复位重测通过。')
    h(d,'配置文件与硬件依据','板级下载文件为 Qproject/output_files_board/board_selftest.sof，采用 USB Blaster 和 JTAG 临时配置，掉电后配置丢失。管脚依据征途引脚绑定映射表 Sheet1 第 3、71、79 至 82 行，以及征途 Pro 开发板硬件规格书 V1.0.1；Pro 手册第 21 页给出用户 LED D6 至 D9 与四路 LED 信号的关系。')

def implementation(r):
    d=new('AXI4 Stream FIR DSP 加速器实现与验证报告','本报告说明各 RTL 模块的内部结构、关键控制条件和定点运算实现，给出功能仿真、布局布线及板级自检结果，并记录设计问题和对应的修正措施。系统基于 EP4CE10F17C8，采用并行乘加流水、流接口反压和影子系数原子更新。')
    d.add_heading('1 模块组成与实现约定',1)
    table(d,['文件','职责','主要状态'],[
        ['accelerator.v','顶层 数据通道与系数排空提交','pause_settled 及暂停确认'],
        ['fir_pipeline.v','历史缓存 系数表 并行乘加流水','active shadow history tree valid_pipe'],
        ['fixed_quantizer.v','舍入 移位 饱和或回绕','纯组合 无内部状态'],
        ['axis_fifo.v','单时钟浅 FIFO','wr_ptr rd_ptr count'],
        ['axis_async_fifo.v','异步浅 FIFO','二进制 Gray 指针 full empty'],
        ['axis_register.v','单拍弹性接口寄存器','payload valid'],
        ['reset_sync.v','异步复位 同步释放','两级 sync_ff']], [1.7,2.65,2.5])
    p(d,'可综合核心使用 Verilog 2001，通过局部 clog2 常量函数和一维数组实现参数化结构。RTL 在不启用 SystemVerilog 选项的 ModelSim 编译中通过。测试平台采用 SystemVerilog 支持断言与自检。两个核心配置共享 RTL，通过 USE_ASYNC 选择输入通道。')
    p(d,'tb_accelerator.sv、tb_quantizer.sv、tb_fifo.sv 和 tb_board_selftest.sv 分别验证数据通道、量化、同步缓存及片上自检。golden_model.py 是独立整数参考。board_selftest.v 实现硬件测试控制，board/board_vectors.v 保存激励、参考结果和系数 ROM。')
    references(d)

    page(d,'FIR 乘加流水线实现')
    equation(d)
    h(d,'样本历史更新','history 存放已经被 FIR 接收的历史输入；当前输入用于抽头 0，history[k−1] 用于其余抽头。只有 valid 与 ready 同时成立时更新历史，未被接收的输入与空拍均不影响滤波状态。计算当前乘积和移位历史在同一上升沿使用非阻塞赋值，乘法看到更新前的历史。')
    h(d,'系数和精度','active 参与乘法，shadow 接受运行时写入。提交时 active 复制整个 shadow；写入并不会改变在途样本的系数。默认乘积位宽 32，累加位宽 35，乘积显式符号扩展。增加 ceil(log2(TAPS)) 个保护位可以覆盖所有抽头求和的最坏幅度。')
    h(d,'加法树结构','LEAVES 为不小于 TAPS 的最小二次幂，未使用叶子写零。tree 使用一维数组，前 LEAVES 项是乘积寄存器；后续层按层号乘 LEAVES 加节点号寻址，将上一层两个节点相加并寄存。TAPS 为 8 时有乘积、三层加法和量化输出共五个运算寄存级；TAPS 为 1 时跳过加法层。')
    h(d,'停顿和有效位','advance 由末级有效位及下游 ready 决定。若末级存在未接收输出，整条运算流水暂停；所有有效位、TLAST、树节点和历史一致冻结。末级为空时仍可推进流水以填满空位。全局使能减少各级独立控制复杂度，但反压控制路径可能成为时序瓶颈。')
    h(d,'帧与延迟定义','FRAME_RESET 为 0 时 TLAST 只标记帧末，历史连续；为 1 时末样本仍使用当前帧历史，其接收后清空历史，不输出额外尾部。无反压的同步顶层从输入握手到输出握手延迟为 ceil(log2(TAPS)) 加 6 拍，默认为 9 拍。功能仿真采用 10 ns 核心周期，测得 90 ns；硬件性能配置的工作周期为 13.333 ns。异步首样本延迟另受时钟相位和指针同步影响。')

    page(d,'定点量化与负数舍入')
    h(d,'可选模式','ROUND_MODE 为 0 时直接算术右移，相当于朝负无穷取整。模式 1 为最近值，中点远离零；模式 2 为最近偶数。SATURATE 为 1 时夹紧到 OUT_WIDTH 的有符号范围，为 0 时取结果低 OUT_WIDTH 位并按补码解释。必须先扩展和舍入，再检查溢出，最后截到输出宽度。')
    table(d,['SHIFT为1 的输入','算术截位','中点远离零','最近偶数'],[[1,0,1,0],[3,1,2,2],[5,2,3,2],[-1,-1,-1,0],[-3,-2,-2,-2],[-5,-3,-3,-2]],[2.05,1.6,1.6,1.6])
    h(d,'实现与优化','din 先符号扩展一位。移位数大于零时，模式 1 的正数偏置为半个量化单位，负数为该值减 1；模式 2 的偏置为半个单位减 1 再加移位后结果的最低位。将有符号输入加偏置后算术右移，等价实现了两种最近舍入。这避免了绝对值、取负和取余形成的长组合路径。SHIFT 为 0 时偏置为零。')
    h(d,'溢出边界','默认输出范围为负 32768 至正 32767。默认 8 抽头全部系数为负 32768，所有样本为负 32768 时，完整累加为正 8589934592；右移 15 位得到 262144，应饱和为 32767。最小负数必须在扩展位宽中参与处理，避免同宽取绝对值溢出。')
    h(d,'验证独立性','Python 使用不限位宽的整数、divmod 和数学范围夹紧，避免照抄 RTL 的位宽截取和舍入偏置实现。量化单元测试另外使用整数商与余数定义，对 12 位输入的全部 4096 个值、六种舍入 饱和组合以及五种 SHIFT 逐一比较。FIR 测试同时覆盖累加、帧历史和最终量化。')

    page(d,'同步 FIFO 与弹性寄存器')
    h(d,'单时钟 FIFO','写指针与读指针只在相应握手时自增，溢出自然回绕。count 区分空满；计数为零不产生 valid。未满时允许写；已满但同拍读出时也允许写入替换。push 与 pop 同时发生时 count 保持，仅单独写入或读出时增减。FIFO 对外先展示当前头部，再在 pop 上升沿推进读指针。')
    h(d,'满时同拍读写','满 FIFO 的写位置是正在消费的最旧位置。时钟边沿下游采样旧数据，写端以非阻塞赋值写入新数据，读指针推进后展示下一项。专用用例覆盖满、满时同拍替换、回绕、顺序及空条件。')
    h(d,'弹性寄存器','valid 表示单个寄存器位置有未消费数据。为空或下游 ready 为 1 时，s_ready 为 1；此时更新 valid，若输入 valid 同时为 1 才更新 payload。已有数据且下游不 ready 时，valid 和 payload 保持。该模块可以每拍替换输出，稳态吞吐为每拍一个样本。')
    h(d,'接口寄存器与存储映射','FIFO 读出与外部接口直接相连会形成较长数据路径。输入和输出 FIFO 后分别设置 payload 寄存器，浅存储声明 ramstyle 为 logic，使默认 16 深度 FIFO 映射为逻辑寄存器，以额外缓存和延迟缩短组合读出路径。')
    h(d,'容量与控制取舍','默认每个 FIFO 16 个位置，输入和输出弹性寄存器各 1 个位置，运算流水 5 个位置。FIFO 可以缓冲有限突发，不能长期消除输入速率高于输出速率的问题。当前 ready 组合链仍存在，性能报告中可以看见这条控制路径；更高频率可用带 skid buffer 的寄存 ready 结构。')

    page(d,'异步 FIFO 与原子配置提交')
    h(d,'Gray 指针','ADDR_WIDTH 加 1 位二进制指针保留回绕信息，每次实际读写加 1，再转换为 Gray。读指针同步到写域，写指针同步到读域。空判断使用下一个读 Gray 指针与同步写 Gray 指针比较；满判断将同步读 Gray 指针的最高两位反转后与下一个写 Gray 比较。两级同步器带有 Quartus 的同步器识别属性。')
    h(d,'满空保守性','跨域指针有两拍延迟，因此状态可以保守地暂时报告空或满。读端不会访问尚未通过写指针同步证明有效的数据，写端不会覆盖仍未确认读出的内容。Gray 总线不能仅靠 RTL 编码保证物理到达次序，需要对位间延迟实施约束。最终 max skew 报告的实际最坏值为约 0.402 ns，上限为 10 ns。')
    h(d,'暂停确认的竞态','源域最后一次写入和暂停确认可能在同一源时钟边沿发生。若核心只看确认已经返回而不考虑 empty 寄存器使用上一拍同步指针，就可能过早判空。pause_settled 在确认后额外延时，再与整个通道的空闲条件共同决定 cfg_ready，从而覆盖该竞态。')
    h(d,'更新事务的边界','系数写入可与数据运行并行，只改变 shadow。cfg_quiesce 停止新的源握手，但所有已接收数据必须继续消费。提交只在 cfg_ready 时生效；清空 history 是定义好的切换边界，不将上一套滤波状态继承到新滤波器。下游长期阻塞会使等待排空持续，这是协议行为而非死锁缺陷。')
    h(d,'跨域验证范围','异步功能测试分别使用 14 ns 与 6 ns 源周期，核心周期为 10 ns，覆盖源端慢与快、长反压、指针回绕和共同复位。断言检查本域 Gray 每次至多变化一位。数字仿真不生成模拟亚稳态，跨域可靠性还依赖同步器结构、芯片参数和物理布线。双时钟版本的验证不包含实物试验。')

    page(d,'功能验证方案与结果')
    table(d,['用例组','重点','数量'],[
        ['同步默认及异步默认','16位 8抽头 3次系数编程','2'],
        ['非二次幂抽头','3抽头 5抽头 小位宽 跨帧清历史','4'],
        ['单抽头与边界舍入','SHIFT为0 中点偶数 中点远离零','3'],
        ['16抽头及快源 CDC','大加法树 源快于核心','2'],
        ['量化穷举','SHIFT为0 1 3 7 11；每组24576比对','5'],
        ['同步 FIFO 单元','满状态同拍读写与指针回绕','1'],
        ['板级自检','黄金 ROM 复位重跑 心跳 故障注入','1']], [2.0,3.6,1.25])
    p(d,f"总计 {r['simulation_tests']} 个用例均 PASS。11 个 FIR 用例各 1536 个样本、三套系数，共 {r['fir_comparisons']} 个输出比对；量化比对 {r['quantizer_comparisons']} 项。向量含冲激、正负极值、交替符号、连续常量、随机数据和周期性 TLAST。")
    h(d,'测试平台行为','源数据在下降沿驱动，在上升沿等待握手，停顿时保持 valid 和 payload。输出消费者首先持续 ready 验证吞吐，然后引入随机与长阻塞。scoreboard 在输出握手时比较数值和 TLAST，接受数与输出数必须一致。停顿稳定性、同步首样本延迟和 Gray 变化均有断言。')
    h(d,'复位与错误注入','正常比较前先令输出阻塞并填满 FIFO，然后共同复位，确认旧数据不会泄漏。每套系数更新前等待排空，流中途还执行暂停排空并恢复。覆盖非法地址、未暂停提交以及同拍 write 和 commit 拒绝；这些行为不会破坏当前有效系数的数值参考。')
    h(d,'证据位置','compile_lint.log 同时记录 RTL 的 Verilog 模式编译和测试平台的 sv 模式编译，均无错误或警告。每个 FIR 用例日志记录接受 输出 停顿 连续输出计数及首样本延迟；板级日志记录两次共 512 个比对及故障检测。WLF 保留信号，summary.json 和 verification_summary.json 提供机器可读摘要。')

    page(d,'设计问题与解决方案')
    h(d,'异步用例未覆盖输入反压','第一轮异步默认测试的输入比核心慢，随机输出 ready 平均吞吐仍大于输入吞吐，导致 input_stalls 为零。数值通过不能说明反压覆盖完整。增加周期性 80 拍阻塞和快源 CDC 用例，最终异步默认输入停顿 805 拍，快源用例为 3751 拍。')
    h(d,'存储读出导致初版时序失败','初版在 100 MHz 下最差 setup 裕量为负 5.080 ns，Fmax 为 66.31 MHz，关键路径由输出 FIFO 的 RAM 地址寄存器到 m_axis_tlast。增加输入 输出接口寄存器、浅 FIFO 使用逻辑存储并简化量化运算后，控制链成为主要限制；100 MHz 尝试仍未闭合，因此最终工作约束设为 75 MHz并完整重编译。')
    h(d,'跨域约束与版本兼容','Quartus 13 不支持 set_max_delay 的 datapath_only 选项，导致原跨域约束无法通过实现。约束改用该版本支持的 set_max_delay、set_min_delay 和 set_max_skew 组合，并避免整体异步时钟组切断屏蔽 Gray 检查。最终 skew 报告匹配 20 条实际路径且无违例，全部约束完整性检查通过。')
    h(d,'配置暂停和指针同步的窗口','设计审查发现暂停 ack 与最后一个 Gray 写指针可能一起返回，而 empty 标志在下一拍才能更新。加入 pause_settled 延时并检查五个数据通道位置全部排空，流中暂停用例要求 cfg_ready 时输出数等于接收数。该问题属于设计审查发现的竞态，未观察到实际硬件故障。')
    h(d,'字节宽度与 AXI 总线','数值参数允许 12 位输入、10 位输出，但 AXI 数据总线必须按字节对齐。顶层对 TDATA 向上补齐并在输出符号扩展，数值核心继续使用原精度；非整字节参数用例直接验证了符号扩展后的总线值。')
    h(d,'复位输入的约束声明','异步复位外部路径采用时序例外，但端口仍应明确声明输入延迟。为 aresetn 补入最大和最小均为零的输入延迟，使 check_timing 能完整识别接口条件。最终 no_input_delay 为零，复位同步结构的本域恢复和移除检查通过。')

    page(d,'最终时序资源与关键路径')
    table(d,['配置','核心 Fmax','LE','寄存器','9位乘法资源','Setup ns'],measured_rows(r),[1.45,1.12,.6,.75,1.25,1.08])
    for rev,v in r['revisions'].items():
        z=v['slow_85C_slack_ns']
        p(d,f"{rev} 在慢模型 1.2 V 85 ℃ 的 setup 为 {z['setup']} ns，hold 为 {z['hold']} ns，recovery 为 {z['recovery']} ns，removal 为 {z['removal']} ns。完整 STA 同时检查慢模型 0 ℃ 和快模型 0 ℃，工作约束均满足。")
    h(d,'Fmax 与工作频率','Fmax 是当前布局和约束下的同钟路径估计，不覆盖异步跨域路径。核心性能配置工作频率为 75 MHz，异步源为 50 MHz，并结合 setup、hold、recovery、removal、pulse width 和约束完整性判断是否满足要求。参数、约束、管脚或布局改变后应重新评估。')
    # Extract actual endpoints directly from final setup report.
    for rev in ['accelerator','accelerator_cdc']:
        lines=(ROOT/'Qproject/reports'/f'{rev}_setup.txt').read_text().splitlines()
        data=[s for s in lines if s.strip().startswith(';') and 'core_clk' in s]
        if data:
            cols=[s.strip() for s in data[0].split(';')][1:-1]
            p(d,f"{rev} 最差 setup 路径：{cols[1]} 至 {cols[2]}。其 slack 为 {cols[0]} ns，说明性能限制包含控制信号传播和外部时序预算，而不只是乘法器运算。")
    h(d,'约束与资源复核','两种配置 no_clock、multiple_clock、no_input_delay、no_output_delay 均为零，report_ucp 的未约束输入 输出路径为零。Gray skew 最差裕量为 9.598 ns。当前使用 0 个 M9K 存储位与 16 个 9 位乘法元素，逻辑使用率约 14%，满足所选芯片资源容量。')
    p(d,'两个核心配置用于性能评估，外部 AXI 管脚自动分配。板级配置已绑定全部六个端口，并采用实际 50 MHz 晶振时钟约束；单时钟实物自检和复位重测通过，结果记录见板级验证章节。')
    board_page(d,r)

    questions=[
        ('为什么选择 FIR','有限冲激响应结构简单，没有反馈环路，定点误差不会沿反馈累积；每个样本都能独立展开并行乘加，适合 FPGA 流水线。'),
        ('TAPS 和阶数有什么区别','TAPS 为抽头个数，阶数是 TAPS 减 1。默认 8 抽头对应 7 阶；增加抽头需要更多乘法资源和加法树级数。'),
        ('为什么每拍能输出一个样本','各抽头乘法并行，每级树节点寄存，相邻样本在不同流水级同时运行。延迟增加但启动间隔为 1；下游阻塞时吞吐由实际握手决定。'),
        ('valid 为 1 时能否改变数据','只在 ready 同时为 1 的握手之后可以换下一份数据。valid 为 1 且 ready 为 0 时必须保持数据、TLAST 和 valid，直到被消费。'),
        ('什么时候移位历史样本','只在 FIR 输入被接收时移位。源端 valid 但未 ready，或流水线暂停，都不能更新历史，否则同一数据会被重复计算。'),
        ('为什么乘积和累加要显式有符号','无符号扩展会将负数变成大正数。乘积用有符号样本和系数，求和前符号扩展到完整累加位宽，避免表达式隐含位宽造成溢出。'),
        ('TLAST 是否意味着复位滤波器','由 FRAME_RESET 决定。默认只传递包边界，跨包历史继续；选择清历史时，末样本仍使用此前历史，然后准备下一帧零初态。'),
        ('FIR 是否自动输出尾部','不自动补零，输入与输出握手数量一致。若应用需要完整卷积尾部，源端应主动追加 TAPS 减 1 个零样本并安排 TLAST。'),
        ('为什么不能运行中直接改 active 系数','不同抽头或不同在途样本可能使用混合的新旧系数，破坏黄金模型定义。影子写入与排空后原子提交给出明确切换边界。'),
        ('为何提交后清空历史','历史样本代表旧滤波状态。本项目定义新系数启用时从零历史开始，便于软件使用和验证；若要无缝保留状态，需要另行定义更新语义。'),
        ('算术右移是否等于朝零截断','不是。负数算术右移相当于向负无穷取整，例如负 3 右移一位得到负 2，而朝零取整会得到负 1。'),
        ('最近偶数舍入为什么有用','中点始终偏向同一方向会带来偏置。最近偶数使中点落在偶数结果上，在许多数值分布中可以降低累计舍入偏差。'),
        ('饱和和回绕何时选择','饱和避免幅度溢出后突然翻转符号，常用于 DSP 输出。回绕对应固定宽度补码运算，适合明确允许模数行为的应用。'),
        ('为什么 Gray 指针仍需约束','逻辑上一次只变一位，但不同位经过不等布线延迟，多个连续状态仍可能在采样端重叠。位间 skew 必须小于发起更新的最短周期。'),
        ('同步器能否保证绝不亚稳','不能。同步器提供恢复时间并显著降低亚稳传播概率，可靠性与芯片参数、时钟 数据频率及布线有关；数字仿真不能直接证明模拟可靠性。'),
        ('能否单独复位 FIFO 一侧','当前设计要求共同复位。单侧清指针会使另一侧仍使用旧回绕状态，存在重读、覆盖或空满误判风险。'),
        ('cfg_ready 为什么一直不出现','先确认 cfg_quiesce 已保持、两个时钟在运行、下游 ready 能消费，以及源端停止条件已同步。任何在途数据未排空时都不允许提交。'),
        ('Fmax 是否就是可上板的频率','不是。Fmax 是给定约束和布局下的同钟路径报告。板级可用频率还需外部接口、管脚、电压、时钟质量和全部时序角落满足要求。'),
        ('如何证明设计没有丢数据','比较输入与输出握手数量，对每次输出的数值和 TLAST 按序比对，再覆盖长反压、时钟速率变化、FIFO 回绕和复位。当前回归均实现了这些检查。'),
        ('怎样继续提高频率','先看真实最差路径。当前可考虑用 skid buffer 寄存 ready、减少配置控制输出的组合层次，并结合芯片资源重新流水化；每次改动重新回归与 STA。'),
    ]
    for group in range(3):
        page(d,['附录 技术问答 数据通道','附录 技术问答 数值与跨域','附录 技术问答 配置与验证'][group])
        for idx,(question,answer) in enumerate(questions[group*7:(group+1)*7],group*7+1):
            para=d.add_paragraph();run=para.add_run(f'{idx}  {question}');run.bold=True
            para.paragraph_format.space_before=Pt(8);para.paragraph_format.space_after=Pt(4)
            para.paragraph_format.keep_with_next=True
            answer_para=d.add_paragraph(answer)
            answer_para.add_run().add_break()
            answer_para.paragraph_format.keep_together=True
            answer_para.paragraph_format.space_after=Pt(8)
    d.save(ROOT/'doc/02_模块实现与验证报告.docx')

if __name__=='__main__':
    r=collect();planning(r);implementation(r)
    print('Created both Word documents under doc')
